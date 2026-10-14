from __future__ import annotations

import csv
import base64
from copy import deepcopy
import json
import os
import sys
from dataclasses import dataclass
from datetime import datetime, timezone
from hashlib import sha256
from io import BytesIO, StringIO
from pathlib import Path
from uuid import uuid4

from openpyxl import load_workbook

from app.schemas.experiments import (
    ExperimentAlignmentAssetSummary,
    ExperimentAlignmentPreviewResponse,
    ExperimentBatchFileUploadRequest,
    ExperimentBatchFileUploadResponse,
    ExperimentAssetCleaningPreview,
    ExperimentConfirmFieldMappingsRequest,
    ExperimentConfirmedFieldMappingSet,
    ExperimentCleaningPreviewResponse,
    ExperimentCleaningQualitySummary,
    ExperimentFieldDictionary,
    ExperimentFieldDictionaryEntry,
    ExperimentFieldDictionaryUpsertRequest,
    ExperimentFeatureAnomalyHint,
    ExperimentFeaturePreviewResponse,
    ExperimentFeatureSummary,
    ExperimentDatasetSnapshot,
    ExperimentModelingPlanPreview,
    ExperimentTrainingAgentStage,
    ExperimentResearchBriefSummary,
    ExperimentSimulationCandidatePlan,
    ExperimentSimulationInputSummary,
    ExperimentSimulationRange,
    ExperimentTrainingResultMetrics,
    ExperimentTrainingResultSummary,
    ExperimentTrainingTaskPreview,
    ExperimentFileAuditEvent,
    ExperimentFileAuditStageSummary,
    ExperimentFileAuditSummary,
    ExperimentFileAuditTimeline,
    ExperimentFieldMappingSuggestion,
    ExperimentFieldMappingSuggestionResponse,
    ExperimentAssetFieldMappingSuggestions,
    ExperimentFileAsset,
    ExperimentFileDeleteResult,
    ExperimentTextFileUploadRequest,
)
from app.schema_recognition.contracts import SchemaAssetContext, SchemaFieldContext, SchemaRecognitionContext
from app.services.experiment_data_store import SqliteExperimentDataStore


@dataclass
class StoredExperimentFileAsset:
    asset_id: str
    experiment_id: str
    original_name: str
    detected_format: str
    parse_status: str
    row_count: int
    column_names: list[str]
    rows: list[dict[str, str]]
    content_hash: str
    derived_from_asset_id: str | None = None
    source_conversion_note: str | None = None


class ExperimentFileService:
    def __init__(self, store: SqliteExperimentDataStore | None = None) -> None:
        self.store = store
        self._assets_by_hash: dict[tuple[str, str], StoredExperimentFileAsset] = {}
        self._audits_by_experiment: dict[str, list[ExperimentFileAuditEvent]] = {}
        self._confirmed_field_mappings_by_experiment: dict[
            str,
            ExperimentConfirmedFieldMappingSet,
        ] = {}
        self._field_dictionary_by_experiment: dict[str, ExperimentFieldDictionary] = {}
        self._dataset_snapshots_by_experiment: dict[str, ExperimentDatasetSnapshot] = {}
        self._modeling_plans_by_experiment: dict[str, ExperimentModelingPlanPreview] = {}
        self._training_tasks_by_experiment: dict[str, ExperimentTrainingTaskPreview] = {}
        self._training_results_by_experiment: dict[str, ExperimentTrainingResultSummary] = {}
        self._research_briefs_by_experiment: dict[str, ExperimentResearchBriefSummary] = {}
        self._simulation_inputs_by_experiment: dict[str, ExperimentSimulationInputSummary] = {}
        self._restore_persistent_state()

    def _restore_persistent_state(self) -> None:
        if self.store is None:
            return
        for payload in self.store.list("file_asset"):
            asset = StoredExperimentFileAsset(**payload)
            self._assets_by_hash[(asset.experiment_id, asset.content_hash)] = asset
        for payload in self.store.list("file_audit"):
            event = ExperimentFileAuditEvent.model_validate(payload)
            self._audits_by_experiment.setdefault(event.experimentId, []).append(event)
        for payload in self.store.list("field_mapping_set"):
            item = ExperimentConfirmedFieldMappingSet.model_validate(payload)
            self._confirmed_field_mappings_by_experiment[item.experimentId] = item
        for payload in self.store.list("dataset_snapshot"):
            item = ExperimentDatasetSnapshot.model_validate(payload)
            self._dataset_snapshots_by_experiment[item.experimentId] = item

    def _persist_asset(self, asset: StoredExperimentFileAsset, content: bytes) -> None:
        if self.store is None:
            return
        self.store.put_bytes(content, asset.content_hash)
        self.store.save(
            "file_asset", asset.asset_id, asset.experiment_id,
            {
                "asset_id": asset.asset_id, "experiment_id": asset.experiment_id,
                "original_name": asset.original_name, "detected_format": asset.detected_format,
                "parse_status": asset.parse_status, "row_count": asset.row_count,
                "column_names": asset.column_names, "rows": asset.rows,
                "content_hash": asset.content_hash, "derived_from_asset_id": asset.derived_from_asset_id,
                "source_conversion_note": asset.source_conversion_note,
            },
        )

    def save_text_file(
        self,
        experiment_id: str,
        request: ExperimentTextFileUploadRequest,
    ) -> ExperimentFileAsset:
        raw_content = _decode_upload(request)
        content_hash = sha256(raw_content).hexdigest()
        source_asset_id = self._resolve_source_asset(experiment_id, request.sourceAssetId)
        source_conversion_note = _normalize_source_conversion_note(request.sourceConversionNote, source_asset_id)
        duplicate_key = (experiment_id, content_hash)
        existing = self._assets_by_hash.get(duplicate_key)
        if existing is not None:
            response = _to_response(
                existing,
                duplicate=True,
                duplicate_of_asset_id=existing.asset_id,
            )
            self._write_audit_event(experiment_id, "upload_file", [response])
            return response

        detected_format = _detect_format(request.fileName)
        column_names, row_count, rows, parse_status = _parse_upload(raw_content, detected_format)
        asset = StoredExperimentFileAsset(
            asset_id=f"file-{uuid4()}",
            experiment_id=experiment_id,
            original_name=Path(request.fileName).name,
            detected_format=detected_format,
            parse_status=parse_status,
            row_count=row_count,
            column_names=column_names,
            rows=rows,
            content_hash=content_hash,
            derived_from_asset_id=source_asset_id,
            source_conversion_note=source_conversion_note,
        )
        self._assets_by_hash[duplicate_key] = asset
        self._persist_asset(asset, raw_content)
        response = _to_response(asset, duplicate=False, duplicate_of_asset_id=None)
        self._write_audit_event(experiment_id, "upload_file", [response])
        return response

    def save_text_files_batch(
        self,
        experiment_id: str,
        request: ExperimentBatchFileUploadRequest,
    ) -> ExperimentBatchFileUploadResponse:
        assets = [
            self._save_text_file_without_audit(experiment_id=experiment_id, request=file)
            for file in request.files
        ]
        self._write_audit_event(experiment_id, "upload_files_batch", assets)
        return ExperimentBatchFileUploadResponse(
            experimentId=experiment_id,
            assetCount=len(assets),
            assets=assets,
        )

    def list_audit_events(self, experiment_id: str) -> ExperimentFileAuditTimeline:
        events = self._audits_by_experiment.get(experiment_id, [])
        return ExperimentFileAuditTimeline(
            experimentId=experiment_id,
            eventCount=len(events),
            events=events,
        )

    def list_file_assets(self, experiment_id: str) -> list[ExperimentFileAsset]:
        assets = [
            asset
            for (stored_experiment_id, _content_hash), asset in self._assets_by_hash.items()
            if stored_experiment_id == experiment_id
        ]
        return [
            _to_response(asset, duplicate=False, duplicate_of_asset_id=None)
            for asset in assets
        ]

    def get_stored_assets(self, experiment_id: str) -> list[StoredExperimentFileAsset]:
        return deepcopy([
            asset for (stored_experiment_id, _), asset in self._assets_by_hash.items()
            if stored_experiment_id == experiment_id
        ])

    def get_analysis_stored_assets(self, experiment_id: str) -> list[StoredExperimentFileAsset]:
        """Only parsed assets may enter field mapping, cleaning, and model training."""
        return [
            asset for asset in self.get_stored_assets(experiment_id)
            if asset.parse_status == "parsed"
        ]

    def get_analysis_file_assets(self, experiment_id: str) -> list[ExperimentFileAsset]:
        return [
            _to_response(asset, duplicate=False, duplicate_of_asset_id=None)
            for asset in self.get_analysis_stored_assets(experiment_id)
        ]

    def get_source_evidence_stored_assets(
        self,
        experiment_id: str,
        analysis_assets: list[StoredExperimentFileAsset] | None = None,
    ) -> list[StoredExperimentFileAsset]:
        analysis_assets = analysis_assets if analysis_assets is not None else self.get_analysis_stored_assets(experiment_id)
        source_ids = {asset.derived_from_asset_id for asset in analysis_assets if asset.derived_from_asset_id}
        return [asset for asset in self.get_stored_assets(experiment_id) if asset.asset_id in source_ids]

    def _resolve_source_asset(self, experiment_id: str, source_asset_id: str | None) -> str | None:
        if source_asset_id is None:
            return None
        source = next((item for item in self._assets_by_hash.values() if item.asset_id == source_asset_id), None)
        if source is None or source.experiment_id != experiment_id:
            raise ValueError("source raw asset was not found in this experiment")
        if source.parse_status != "awaiting_adapter":
            raise ValueError("source asset must be awaiting device adaptation")
        return source.asset_id

    def build_schema_recognition_context(self, experiment_id: str, sample_limit: int = 3) -> SchemaRecognitionContext:
        limit = max(1, min(sample_limit, 5))
        assets = self.get_analysis_stored_assets(experiment_id)
        return SchemaRecognitionContext(
            experimentId=experiment_id,
            assets=[SchemaAssetContext(
                assetId=asset.asset_id,
                originalName=asset.original_name,
                detectedFormat=asset.detected_format,
                fields=[SchemaFieldContext(
                    originalColumn=column,
                    sampleValues=[row[column] for row in asset.rows[:limit] if row.get(column) not in {None, ""}],
                ) for column in asset.column_names],
            ) for asset in assets],
        )

    def suggest_field_mappings(self, experiment_id: str) -> ExperimentFieldMappingSuggestionResponse:
        assets = self.get_analysis_stored_assets(experiment_id)
        dictionary = self._field_dictionary_by_experiment.get(
            experiment_id,
            ExperimentFieldDictionary(experimentId=experiment_id, entryCount=0, entries=[]),
        )
        return ExperimentFieldMappingSuggestionResponse(
            experimentId=experiment_id,
            assetCount=len(assets),
            assets=[
                ExperimentAssetFieldMappingSuggestions(
                    assetId=asset.asset_id,
                    originalName=asset.original_name,
                    suggestions=[
                        _suggest_field_mapping(column, dictionary.entries)
                        for column in asset.column_names
                    ],
                )
                for asset in assets
            ],
        )

    def confirm_field_mappings(
        self,
        experiment_id: str,
        request: ExperimentConfirmFieldMappingsRequest,
    ) -> ExperimentConfirmedFieldMappingSet:
        mapping_set = ExperimentConfirmedFieldMappingSet(
            experimentId=experiment_id,
            mappingCount=len(request.mappings),
            mappings=request.mappings,
        )
        self._confirmed_field_mappings_by_experiment[experiment_id] = mapping_set
        if self.store is not None:
            self.store.save(
                "field_mapping_set", experiment_id, experiment_id,
                mapping_set.model_dump(mode="json"),
            )
        return mapping_set

    def get_confirmed_field_mappings(
        self,
        experiment_id: str,
    ) -> ExperimentConfirmedFieldMappingSet:
        return self._confirmed_field_mappings_by_experiment.get(
            experiment_id,
            ExperimentConfirmedFieldMappingSet(
                experimentId=experiment_id,
                mappingCount=0,
                mappings=[],
            ),
        )

    def upsert_field_dictionary(
        self,
        experiment_id: str,
        request: ExperimentFieldDictionaryUpsertRequest,
    ) -> ExperimentFieldDictionary:
        existing = self._field_dictionary_by_experiment.get(
            experiment_id,
            ExperimentFieldDictionary(experimentId=experiment_id, entryCount=0, entries=[]),
        )
        entries_by_source = {
            _normalize_column_name(entry.sourceColumn): entry
            for entry in existing.entries
        }
        for entry in request.entries:
            entries_by_source[_normalize_column_name(entry.sourceColumn)] = entry
        dictionary = ExperimentFieldDictionary(
            experimentId=experiment_id,
            entryCount=len(entries_by_source),
            entries=list(entries_by_source.values()),
        )
        self._field_dictionary_by_experiment[experiment_id] = dictionary
        return dictionary

    def get_field_dictionary(self, experiment_id: str) -> ExperimentFieldDictionary:
        return self._field_dictionary_by_experiment.get(
            experiment_id,
            ExperimentFieldDictionary(experimentId=experiment_id, entryCount=0, entries=[]),
        )

    def preview_cleaning(self, experiment_id: str) -> ExperimentCleaningPreviewResponse:
        assets = self.get_analysis_stored_assets(experiment_id)
        confirmed_mappings = self.get_confirmed_field_mappings(experiment_id)
        mappings_by_asset_and_column = {
            (mapping.assetId, mapping.originalColumn): mapping
            for mapping in confirmed_mappings.mappings
        }
        previews = [
            _build_cleaning_preview(asset, mappings_by_asset_and_column)
            for asset in assets
        ]
        return ExperimentCleaningPreviewResponse(
            experimentId=experiment_id,
            assetCount=len(previews),
            assets=previews,
        )

    def preview_alignment(self, experiment_id: str) -> ExperimentAlignmentPreviewResponse:
        cleaning_preview = self.preview_cleaning(experiment_id)
        timestamps = sorted(
            {
                row.get("timestamp")
                for asset in cleaning_preview.assets
                for row in asset.cleanedRows
                if row.get("timestamp")
            }
        )
        fields = sorted(
            {
                field
                for asset in cleaning_preview.assets
                for row in asset.cleanedRows
                for field in row
                if field != "timestamp"
            }
        )
        value_by_timestamp_and_field: dict[tuple[str, str], str | float | None] = {}
        asset_summaries: list[ExperimentAlignmentAssetSummary] = []
        for asset in cleaning_preview.assets:
            asset_fields = sorted(
                {
                    field
                    for row in asset.cleanedRows
                    for field in row
                    if field != "timestamp"
                }
            )
            asset_timestamps = {
                row.get("timestamp")
                for row in asset.cleanedRows
                if row.get("timestamp")
            }
            asset_summaries.append(
                ExperimentAlignmentAssetSummary(
                    assetId=asset.assetId,
                    originalName=asset.originalName,
                    timestampCount=len(asset_timestamps),
                    fields=asset_fields,
                )
            )
            for row in asset.cleanedRows:
                timestamp = row.get("timestamp")
                if not isinstance(timestamp, str):
                    continue
                for field in asset_fields:
                    value_by_timestamp_and_field[(timestamp, field)] = row.get(field)

        aligned_rows: list[dict[str, str | float | None]] = []
        alignment_logs: list[str] = []
        missing_cell_count = 0
        for timestamp in timestamps:
            aligned_row: dict[str, str | float | None] = {"timestamp": timestamp}
            for field in fields:
                value = value_by_timestamp_and_field.get((timestamp, field))
                aligned_row[field] = value
                if value is None:
                    missing_cell_count += 1
                    alignment_logs.append(f"Missing aligned value: {field} at {timestamp}")
            aligned_rows.append(aligned_row)

        return ExperimentAlignmentPreviewResponse(
            experimentId=experiment_id,
            assetCount=cleaning_preview.assetCount,
            timestampCount=len(timestamps),
            missingCellCount=missing_cell_count,
            assetSummaries=asset_summaries,
            alignedRows=aligned_rows,
            alignmentLogs=alignment_logs,
        )

    def preview_features(self, experiment_id: str) -> ExperimentFeaturePreviewResponse:
        alignment_preview = self.preview_alignment(experiment_id)
        fields = sorted(
            {
                field
                for row in alignment_preview.alignedRows
                for field, value in row.items()
                if field != "timestamp" and isinstance(value, (int, float))
            }
        )
        features: list[ExperimentFeatureSummary] = []
        anomaly_hints: list[ExperimentFeatureAnomalyHint] = []
        for field in fields:
            timestamped_values = [
                (str(row.get("timestamp")), row[field])
                for row in alignment_preview.alignedRows
                if isinstance(row.get(field), (int, float))
            ]
            values = [float(value) for _timestamp, value in timestamped_values]
            if not values:
                continue
            mean = round(sum(values) / len(values), 4)
            first = values[0]
            last = values[-1]
            delta = round(last - first, 4)
            features.append(
                ExperimentFeatureSummary(
                    field=field,
                    count=len(values),
                    min=min(values),
                    max=max(values),
                    mean=mean,
                    first=first,
                    last=last,
                    delta=delta,
                    windowMean=mean,
                )
            )
            anomaly_hints.extend(_detect_feature_anomalies(field, timestamped_values))

        return ExperimentFeaturePreviewResponse(
            experimentId=experiment_id,
            rowCount=len(alignment_preview.alignedRows),
            featureCount=len(features),
            features=features,
            anomalyHints=anomaly_hints,
        )

    def create_dataset_snapshot(
        self,
        experiment_id: str,
        *,
        aligned_rows_override: list[dict[str, str | float | None]] | None = None,
        cleaning_run_id: str | None = None,
        cleaning_output_hash: str | None = None,
    ) -> ExperimentDatasetSnapshot:
        assets = self.get_analysis_stored_assets(experiment_id)
        source_evidence_assets = self.get_source_evidence_stored_assets(experiment_id, assets)
        confirmed_mappings = self.get_confirmed_field_mappings(experiment_id)
        if not assets:
            raise ValueError("at least one parsed experiment file asset is required")
        if confirmed_mappings.mappingCount == 0:
            raise ValueError("at least one human-confirmed field mapping is required")
        live_asset_ids = {asset.asset_id for asset in assets}
        stale_mappings = [
            mapping for mapping in confirmed_mappings.mappings
            if mapping.assetId not in live_asset_ids
        ]
        if stale_mappings:
            raise ValueError("confirmed field mapping references a deleted file asset")
        alignment_preview = self.preview_alignment(experiment_id)
        feature_preview = self.preview_features(experiment_id)
        aligned_rows = (
            aligned_rows_override
            if aligned_rows_override is not None
            else alignment_preview.alignedRows
        )
        numeric_fields = {
            key for row in aligned_rows for key, value in row.items()
            if key != "timestamp" and isinstance(value, (int, float))
        }
        content_payload = {
            "analysisAssets": [
                {"assetId": asset.asset_id, "contentHash": asset.content_hash}
                for asset in sorted(assets, key=lambda item: item.asset_id)
            ],
            "sourceEvidenceAssets": [
                {"assetId": asset.asset_id, "contentHash": asset.content_hash}
                for asset in sorted(source_evidence_assets, key=lambda item: item.asset_id)
            ],
            "mappingCount": confirmed_mappings.mappingCount,
            "alignedRows": aligned_rows,
            "features": [feature.model_dump() for feature in feature_preview.features],
            "anomalyHints": [hint.model_dump() for hint in feature_preview.anomalyHints],
        }
        content_hash = sha256(
            json.dumps(content_payload, sort_keys=True, ensure_ascii=False).encode("utf-8")
        ).hexdigest()
        snapshot = ExperimentDatasetSnapshot(
            snapshotId=f"dataset-{uuid4()}",
            experimentId=experiment_id,
            assetCount=len(assets),
            evidenceAssetCount=len(source_evidence_assets),
            analysisAssetIds=sorted(asset.asset_id for asset in assets),
            sourceEvidenceAssetIds=sorted(asset.asset_id for asset in source_evidence_assets),
            mappingCount=confirmed_mappings.mappingCount,
            alignedRowCount=len(aligned_rows),
            featureCount=len(numeric_fields),
            anomalyCount=len(feature_preview.anomalyHints),
            contentHash=content_hash,
            cleaningRunId=cleaning_run_id,
            cleaningOutputHash=cleaning_output_hash,
            alignedRows=aligned_rows,
            status="modeling_ready",
            createdAt=datetime.now(timezone.utc).isoformat(),
        )
        self._dataset_snapshots_by_experiment[experiment_id] = snapshot
        if self.store is not None:
            self.store.save(
                "dataset_snapshot", experiment_id, experiment_id,
                snapshot.model_dump(mode="json"),
            )
        return snapshot

    def verify_asset_integrity(self, asset_id: str) -> bool:
        asset = next(
            (item for item in self._assets_by_hash.values() if item.asset_id == asset_id),
            None,
        )
        if asset is None:
            raise LookupError("experiment file asset was not found")
        if self.store is None:
            return True
        self.store.read_verified_bytes(asset.content_hash)
        return True

    def get_file_asset(self, experiment_id: str, asset_id: str) -> ExperimentFileAsset | None:
        asset = next(
            (
                item for item in self._assets_by_hash.values()
                if item.experiment_id == experiment_id and item.asset_id == asset_id
            ),
            None,
        )
        return _to_response(asset, False, None) if asset is not None else None

    def delete_file_asset(
        self, experiment_id: str, asset_id: str, operator: str, reason: str,
    ) -> ExperimentFileDeleteResult:
        match = next(
            (
                (key, asset) for key, asset in self._assets_by_hash.items()
                if asset.experiment_id == experiment_id and asset.asset_id == asset_id
            ),
            None,
        )
        if match is None:
            raise LookupError("experiment file asset was not found")
        key, asset = match
        dependent_assets = [
            item.asset_id
            for item in self._assets_by_hash.values()
            if item.experiment_id == experiment_id
            and item.derived_from_asset_id == asset_id
        ]
        if dependent_assets:
            raise ValueError(
                "source asset has derived structured assets and cannot be deleted: "
                + ", ".join(dependent_assets)
            )
        response = _to_response(asset, duplicate=False, duplicate_of_asset_id=None)
        del self._assets_by_hash[key]
        if self.store is not None:
            self.store.soft_delete("file_asset", asset_id)
        event = self._write_audit_event(
            experiment_id, "delete_file", [response],
            detail={"operator": operator, "reason": reason},
        )
        return ExperimentFileDeleteResult(
            assetId=asset_id, experimentId=experiment_id, status="deleted",
            retainedContentHash=asset.content_hash, auditId=event.auditId,
        )

    def get_latest_dataset_snapshot(self, experiment_id: str) -> ExperimentDatasetSnapshot:
        existing = self._dataset_snapshots_by_experiment.get(experiment_id)
        if existing is not None:
            return existing
        return ExperimentDatasetSnapshot(
            snapshotId="",
            experimentId=experiment_id,
            assetCount=0,
            evidenceAssetCount=0,
            analysisAssetIds=[],
            sourceEvidenceAssetIds=[],
            mappingCount=0,
            alignedRowCount=0,
            featureCount=0,
            anomalyCount=0,
            contentHash="",
            status="not_created",
            createdAt="",
        )

    def export_numeric_training_dataset(self, experiment_id: str, snapshot_id: str):
        """Export complete numeric rows only when they still match the immutable snapshot."""
        import numpy as np
        from app.research.sklearn_runner import NumericTrainingDataset

        snapshot = self.get_latest_dataset_snapshot(experiment_id)
        if not snapshot.snapshotId or snapshot.snapshotId != snapshot_id:
            raise ValueError("dataset snapshot is missing or no longer current")
        alignment_rows = (
            snapshot.alignedRows
            if snapshot.alignedRows
            else self.preview_alignment(experiment_id).alignedRows
        )
        plan = self.preview_modeling_plan(experiment_id)
        target = plan.targetField
        if not target:
            raise ValueError("confirmed numeric target field is required")
        numeric_fields = sorted({key for row in alignment_rows for key, value in row.items() if key != "timestamp" and isinstance(value, (int, float))})
        features = [field for field in numeric_fields if field != target]
        if not features:
            raise ValueError("at least one numeric feature distinct from target is required")
        complete = [row for row in alignment_rows if isinstance(row.get(target), (int, float)) and all(isinstance(row.get(field), (int, float)) for field in features)]
        if len(complete) != snapshot.alignedRowCount:
            raise ValueError("aligned rows changed after snapshot; create a new dataset snapshot")
        payload = {"alignedRows": complete, "features": features, "target": target}
        computed = sha256(json.dumps(payload, ensure_ascii=False, sort_keys=True).encode("utf-8")).hexdigest()
        return NumericTrainingDataset(experimentId=experiment_id, snapshotId=snapshot_id,
            datasetFingerprint=f"sha256:{snapshot.contentHash}:{computed}", featureNames=features, targetField=target,
            X=np.asarray([[float(row[field]) for field in features] for row in complete], dtype=float),
            y=np.asarray([float(row[target]) for row in complete], dtype=float))

    def preview_modeling_plan(self, experiment_id: str) -> ExperimentModelingPlanPreview:
        snapshot = self.get_latest_dataset_snapshot(experiment_id)
        feature_preview = self.preview_features(experiment_id)
        confirmed_mappings = self.get_confirmed_field_mappings(experiment_id)

        standard_fields = [
            mapping.standardField
            for mapping in confirmed_mappings.mappings
            if mapping.standardField != "timestamp"
        ]
        unique_fields = list(dict.fromkeys(standard_fields))
        target_field = unique_fields[0] if unique_fields else ""
        candidate_features = [field for field in unique_fields if field != target_field]

        recommended_model_type = (
            "time_series_regression" if snapshot.alignedRowCount > 0 and target_field else "not_ready"
        )
        risk_hints: list[str] = []
        if snapshot.alignedRowCount < 5:
            risk_hints.append(
                f"Only {snapshot.alignedRowCount} aligned rows are available; demo modeling confidence will be low."
            )
        if feature_preview.featureCount <= 1:
            risk_hints.append("Feature diversity is limited; add more process variables before real training.")
        if snapshot.anomalyCount > 0:
            risk_hints.append(
                f"{snapshot.anomalyCount} anomaly hints were detected; review them before promoting this plan."
            )
        if not target_field:
            risk_hints.append("No confirmed target field is available yet.")

        plan = ExperimentModelingPlanPreview(
            experimentId=experiment_id,
            snapshotId=snapshot.snapshotId,
            targetField=target_field,
            candidateFeatures=candidate_features,
            recommendedModelType=recommended_model_type,
            trainSplitRatio=0.8,
            validationSplitRatio=0.2,
            evaluationMetrics=["mae", "rmse", "r2"],
            riskHints=risk_hints,
            status="draft_ready" if snapshot.status == "modeling_ready" and target_field else "not_ready",
        )
        self._modeling_plans_by_experiment[experiment_id] = plan
        return plan

    def create_training_task(self, experiment_id: str) -> ExperimentTrainingTaskPreview:
        plan = self.preview_modeling_plan(experiment_id)
        now = datetime.now(timezone.utc).isoformat()
        task = ExperimentTrainingTaskPreview(
            taskId=f"train-{uuid4()}",
            experimentId=experiment_id,
            snapshotId=plan.snapshotId,
            planTargetField=plan.targetField,
            recommendedModelType=plan.recommendedModelType,
            evaluationMetrics=plan.evaluationMetrics,
            status="queued" if plan.status == "draft_ready" else "not_ready",
            createdAt=now,
            updatedAt=now,
            agentStages=[
                ExperimentTrainingAgentStage(
                    agentName="data_agent",
                    stageLabel="Data validation",
                    status="completed" if plan.status == "draft_ready" else "waiting",
                    detail="Validated dataset snapshot, mappings, and aligned rows.",
                ),
                ExperimentTrainingAgentStage(
                    agentName="modeling_agent",
                    stageLabel="Model fitting",
                    status="queued" if plan.status == "draft_ready" else "waiting",
                    detail="Prepare regression training task and evaluation loop.",
                ),
                ExperimentTrainingAgentStage(
                    agentName="report_agent",
                    stageLabel="Interpretation draft",
                    status="waiting",
                    detail="Prepare result summary, confidence notes, and next-step suggestions.",
                ),
            ],
        )
        self._training_tasks_by_experiment[experiment_id] = task
        return task

    def get_latest_training_task(self, experiment_id: str) -> ExperimentTrainingTaskPreview:
        existing = self._training_tasks_by_experiment.get(experiment_id)
        if existing is None:
            return ExperimentTrainingTaskPreview(
                taskId="",
                experimentId=experiment_id,
                snapshotId="",
                planTargetField="",
                recommendedModelType="not_ready",
                evaluationMetrics=[],
                status="not_created",
                createdAt="",
                updatedAt="",
                agentStages=[],
            )

        progressed_stages = [stage.model_copy(deep=True) for stage in existing.agentStages]
        if existing.status == "queued":
            progressed_stages[1].status = "running"
            progressed = existing.model_copy(
                update={
                    "status": "running",
                    "updatedAt": datetime.now(timezone.utc).isoformat(),
                    "agentStages": progressed_stages,
                }
            )
            self._training_tasks_by_experiment[experiment_id] = progressed
            return progressed

        if existing.status == "running":
            progressed_stages[1].status = "completed"
            progressed_stages[2].status = "completed"
            progressed = existing.model_copy(
                update={
                    "status": "completed",
                    "updatedAt": datetime.now(timezone.utc).isoformat(),
                    "agentStages": progressed_stages,
                }
            )
            self._training_tasks_by_experiment[experiment_id] = progressed
            return progressed

        return existing

    def get_latest_training_result(self, experiment_id: str) -> ExperimentTrainingResultSummary:
        existing = self._training_results_by_experiment.get(experiment_id)
        if existing is not None:
            return existing

        task = self.get_latest_training_task(experiment_id)
        if not task.taskId:
            return ExperimentTrainingResultSummary(
                experimentId=experiment_id,
                taskId="",
                targetField="",
                status="not_ready",
                confidenceLevel="low",
                metrics=ExperimentTrainingResultMetrics(mae=0.0, rmse=0.0, r2=0.0),
                riskNotes=[],
                recommendations=[],
                manualReviewItems=[],
                generatedAt="",
            )

        feature_preview = self.preview_features(experiment_id)
        aligned_rows = self.get_latest_dataset_snapshot(experiment_id).alignedRowCount
        has_anomalies = len(feature_preview.anomalyHints) > 0
        r2 = 0.78 if aligned_rows >= 3 else 0.52
        result = ExperimentTrainingResultSummary(
            experimentId=experiment_id,
            taskId=task.taskId,
            targetField=task.planTargetField,
            status="completed" if task.status == "completed" else "ready",
            confidenceLevel="medium" if has_anomalies or aligned_rows < 5 else "high",
            metrics=ExperimentTrainingResultMetrics(
                mae=0.18,
                rmse=0.24,
                r2=r2,
            ),
            riskNotes=[
                "Current result is generated from a demo training flow rather than a real persisted model.",
                "Dataset size is still limited; confidence should be validated with more batches.",
            ]
            + (
                ["Anomaly hints were detected in the source features and should be reviewed manually."]
                if has_anomalies
                else []
            ),
            recommendations=[
                "Add more aligned experimental batches before promoting this model to process guidance.",
                "Review the highest-variance process variables and confirm whether they reflect real process shifts.",
                "Use this result summary as the basis for the next研发建议书草案。",
            ],
            manualReviewItems=[
                "Confirm target field selection matches the real研发目标。",
                "Check whether anomaly points should be filtered or annotated before real training.",
                "Verify metric thresholds with domain experts before external presentation.",
            ],
            generatedAt=datetime.now(timezone.utc).isoformat(),
        )
        self._training_results_by_experiment[experiment_id] = result
        return result

    def get_latest_research_brief(self, experiment_id: str) -> ExperimentResearchBriefSummary:
        existing = self._research_briefs_by_experiment.get(experiment_id)
        if existing is not None:
            return existing

        result = self.get_latest_training_result(experiment_id)
        if not result.taskId:
            return ExperimentResearchBriefSummary(
                experimentId=experiment_id,
                taskId="",
                title="",
                objective="",
                dataSummary="",
                modelSummary="",
                riskSummary="",
                nextActions=[],
                manualConfirmationItems=[],
                status="not_ready",
                generatedAt="",
            )

        snapshot = self.get_latest_dataset_snapshot(experiment_id)
        feature_preview = self.preview_features(experiment_id)
        brief = ExperimentResearchBriefSummary(
            experimentId=experiment_id,
            taskId=result.taskId,
            title=f"AKK 研发建议书摘要 - {result.targetField or 'process_model'}",
            objective="将当前实验数据整理结果转化为可解释的建模结论，并形成下一轮研发动作建议。",
            dataSummary=(
                f"当前快照包含 {snapshot.assetCount} 份文件资产、{snapshot.alignedRowCount} 条对齐记录、"
                f"{feature_preview.featureCount} 个候选特征，可作为演示版建模输入。"
            ),
            modelSummary=(
                f"模型当前围绕目标字段 {result.targetField} 生成演示结果，"
                f"MAE={result.metrics.mae:.2f}，RMSE={result.metrics.rmse:.2f}，R²={result.metrics.r2:.2f}，"
                f"置信度评级为 {result.confidenceLevel}。"
            ),
            riskSummary="；".join(result.riskNotes[:3]),
            nextActions=[
                "扩充更多批次实验数据，提升训练稳定性。",
                "围绕高波动变量补充复现实验，验证趋势是否稳定。",
                "将当前摘要继续扩展为正式研发建议书或论文草案。",
            ],
            manualConfirmationItems=result.manualReviewItems,
            status="draft_ready" if result.status in ["ready", "completed"] else "not_ready",
            generatedAt=datetime.now(timezone.utc).isoformat(),
        )
        self._research_briefs_by_experiment[experiment_id] = brief
        return brief

    def get_latest_simulation_input(self, experiment_id: str) -> ExperimentSimulationInputSummary:
        existing = self._simulation_inputs_by_experiment.get(experiment_id)
        if existing is not None:
            return existing

        brief = self.get_latest_research_brief(experiment_id)
        if not brief.taskId:
            return ExperimentSimulationInputSummary(
                experimentId=experiment_id,
                targetField="",
                objective="",
                recommendedRanges={},
                candidatePlans=[],
                manualConfirmationItems=[],
                status="not_ready",
                generatedAt="",
            )

        result = self.get_latest_training_result(experiment_id)
        recommended_ranges = {
            "temperature_c": ExperimentSimulationRange(min=36.2, max=37.4, unit="℃"),
            "ph": ExperimentSimulationRange(min=6.7, max=7.1, unit="pH"),
            "duration_hours": ExperimentSimulationRange(min=16.0, max=24.0, unit="h"),
        }
        simulation_input = ExperimentSimulationInputSummary(
            experimentId=experiment_id,
            targetField=result.targetField,
            objective="为下一步产线仿真和条件优化准备可调参的标准化输入对象。",
            recommendedRanges=recommended_ranges,
            candidatePlans=[
                ExperimentSimulationCandidatePlan(
                    planName="稳健基线方案",
                    temperatureC=36.5,
                    ph=6.8,
                    durationHours=20.0,
                    reason="优先保持稳定性，适合作为下一轮复现实验的参考基线。",
                ),
                ExperimentSimulationCandidatePlan(
                    planName="产量拉升方案",
                    temperatureC=37.0,
                    ph=6.9,
                    durationHours=18.0,
                    reason="略提高温度并缩短周期，用于观察产量提升与波动风险。",
                ),
                ExperimentSimulationCandidatePlan(
                    planName="保守验证方案",
                    temperatureC=36.3,
                    ph=7.0,
                    durationHours=22.0,
                    reason="更保守地控制波动，用于验证模型建议是否具有稳健性。",
                ),
            ],
            manualConfirmationItems=[
                "确认温度和 pH 区间是否符合当前设备与菌株耐受范围。",
                "确认发酵时长是否受排产和下游处理节奏限制。",
                "确认是否需要加入搅拌、溶氧等额外参数作为下一步仿真输入。",
            ],
            status="draft_ready" if brief.status == "draft_ready" else "not_ready",
            generatedAt=datetime.now(timezone.utc).isoformat(),
        )
        self._simulation_inputs_by_experiment[experiment_id] = simulation_input
        return simulation_input

    def summarize_audit_events(self, experiment_id: str) -> ExperimentFileAuditSummary:
        events = self._audits_by_experiment.get(experiment_id, [])
        stage_summary_by_stage: dict[str, ExperimentFileAuditStageSummary] = {}
        for event in events:
            if event.stage not in stage_summary_by_stage:
                stage_summary_by_stage[event.stage] = ExperimentFileAuditStageSummary(
                    stage=event.stage,
                    stageLabel=event.stageLabel,
                    eventCount=0,
                    assetCount=0,
                )
            summary = stage_summary_by_stage[event.stage]
            summary.eventCount += 1
            summary.assetCount += len(event.assetIds)

        return ExperimentFileAuditSummary(
            experimentId=experiment_id,
            eventCount=len(events),
            assetCount=sum(len(event.assetIds) for event in events),
            stageSummary=list(stage_summary_by_stage.values()),
            lastEvent=events[-1] if events else None,
        )

    def _save_text_file_without_audit(
        self,
        experiment_id: str,
        request: ExperimentTextFileUploadRequest,
    ) -> ExperimentFileAsset:
        raw_content = _decode_upload(request)
        content_hash = sha256(raw_content).hexdigest()
        source_asset_id = self._resolve_source_asset(experiment_id, request.sourceAssetId)
        source_conversion_note = _normalize_source_conversion_note(request.sourceConversionNote, source_asset_id)
        duplicate_key = (experiment_id, content_hash)
        existing = self._assets_by_hash.get(duplicate_key)
        if existing is not None:
            return _to_response(
                existing,
                duplicate=True,
                duplicate_of_asset_id=existing.asset_id,
            )

        detected_format = _detect_format(request.fileName)
        column_names, row_count, rows, parse_status = _parse_upload(raw_content, detected_format)
        asset = StoredExperimentFileAsset(
            asset_id=f"file-{uuid4()}",
            experiment_id=experiment_id,
            original_name=Path(request.fileName).name,
            detected_format=detected_format,
            parse_status=parse_status,
            row_count=row_count,
            column_names=column_names,
            rows=rows,
            content_hash=content_hash,
            derived_from_asset_id=source_asset_id,
            source_conversion_note=source_conversion_note,
        )
        self._assets_by_hash[duplicate_key] = asset
        self._persist_asset(asset, raw_content)
        return _to_response(asset, duplicate=False, duplicate_of_asset_id=None)

    def _write_audit_event(
        self,
        experiment_id: str,
        action: str,
        assets: list[ExperimentFileAsset],
        detail: dict[str, str] | None = None,
    ) -> ExperimentFileAuditEvent:
        event = ExperimentFileAuditEvent(
            auditId=f"audit-{uuid4()}",
            experimentId=experiment_id,
            action=action,
            stage="file_upload",
            stageLabel="文件上传",
            fileNames=[asset.originalName for asset in assets],
            assetIds=[asset.assetId for asset in assets],
            contentHashes=[asset.contentHash for asset in assets],
            duplicates=[asset.duplicate for asset in assets],
            createdAt=datetime.now(timezone.utc).isoformat(),
            operator=detail.get("operator") if detail else None,
            reason=detail.get("reason") if detail else None,
        )
        self._audits_by_experiment.setdefault(experiment_id, []).append(event)
        if self.store is not None:
            self.store.save(
                "file_audit", event.auditId, experiment_id,
                event.model_dump(mode="json"),
            )
        return event


def _to_response(
    asset: StoredExperimentFileAsset,
    duplicate: bool,
    duplicate_of_asset_id: str | None,
) -> ExperimentFileAsset:
    return ExperimentFileAsset(
        assetId=asset.asset_id,
        experimentId=asset.experiment_id,
        originalName=asset.original_name,
        detectedFormat=asset.detected_format,
        parseStatus=asset.parse_status,
        rowCount=asset.row_count,
        columnNames=asset.column_names,
        contentHash=asset.content_hash,
        duplicate=duplicate,
        duplicateOfAssetId=duplicate_of_asset_id,
        derivedFromAssetId=asset.derived_from_asset_id,
        sourceConversionNote=asset.source_conversion_note,
    )


def _normalize_source_conversion_note(note: str | None, source_asset_id: str | None) -> str | None:
    normalized = note.strip() if note else None
    if normalized and source_asset_id is None:
        raise ValueError("source conversion note requires a source raw asset")
    return normalized


def _content_hash(content: str) -> str:
    return sha256(content.encode("utf-8")).hexdigest()


def _detect_format(file_name: str) -> str:
    suffix = Path(file_name).suffix.lower()
    if suffix == ".csv":
        return "csv"
    if suffix in {".txt", ".tsv"}:
        return "txt"
    if suffix == ".xlsx":
        return "xlsx"
    if suffix == ".xls":
        return "xls"
    if suffix in {".raw", ".dat", ".bin"}:
        return suffix.removeprefix(".")
    return "unknown"


def _decode_upload(request: ExperimentTextFileUploadRequest) -> bytes:
    if request.encoding == "base64":
        try:
            return base64.b64decode(request.content, validate=True)
        except ValueError as exc:
            raise ValueError("invalid base64 file content") from exc
    return request.content.encode("utf-8")


def _parse_upload(content: bytes, detected_format: str) -> tuple[list[str], int, list[dict[str, str]], str]:
    if detected_format in {"csv", "txt"}:
        columns, count, rows = _parse_table_text(content.decode("utf-8-sig"), detected_format)
        return columns, count, rows, "parsed"
    if detected_format == "xlsx":
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        try:
            sheet = next((item for item in workbook.worksheets if item.max_row > 0 and item.max_column > 0), None)
            if sheet is None:
                return [], 0, [], "parsed"
            values = list(sheet.iter_rows(values_only=True))
            if not values:
                return [], 0, [], "parsed"
            columns = [str(value) if value is not None else f"column_{index + 1}" for index, value in enumerate(values[0])]
            rows = [
                {column: "" if index >= len(row) or row[index] is None else str(row[index]) for index, column in enumerate(columns)}
                for row in values[1:] if any(value is not None for value in row)
            ]
            return columns, len(rows), rows, "parsed"
        finally:
            workbook.close()
    return [], 0, [], "awaiting_adapter"


def _parse_table_text(content: str, detected_format: str) -> tuple[list[str], int, list[dict[str, str]]]:
    delimiter = "\t" if detected_format == "txt" and "\t" in content.splitlines()[0] else ","
    reader = csv.reader(StringIO(content), delimiter=delimiter)
    rows = [row for row in reader if row]
    if not rows:
        return [], 0, []
    column_names = rows[0]
    data_rows = [
        {
            column_name: row[index] if index < len(row) else ""
            for index, column_name in enumerate(column_names)
        }
        for row in rows[1:]
    ]
    return column_names, len(data_rows), data_rows


def _build_cleaning_preview(
    asset: StoredExperimentFileAsset,
    mappings_by_asset_and_column: dict[tuple[str, str], object],
) -> ExperimentAssetCleaningPreview:
    cleaned_rows: list[dict[str, str | float | None]] = []
    cleaning_logs: list[str] = []
    missing_value_count = 0
    invalid_numeric_count = 0
    mapped_columns = {
        original_column
        for (asset_id, original_column), _mapping in mappings_by_asset_and_column.items()
        if asset_id == asset.asset_id
    }
    unmapped_column_count = len(
        [column for column in asset.column_names if column not in mapped_columns]
    )

    for row_index, row in enumerate(asset.rows, start=1):
        cleaned_row: dict[str, str | float | None] = {}
        for original_column in asset.column_names:
            mapping = mappings_by_asset_and_column.get((asset.asset_id, original_column))
            if mapping is None:
                continue
            raw_value = row.get(original_column, "")
            if raw_value == "":
                missing_value_count += 1
                cleaned_row[mapping.standardField] = None
                cleaning_logs.append(f"Missing value: {original_column} at row {row_index}")
                continue
            if _is_numeric_standard_field(mapping.standardField):
                try:
                    cleaned_row[mapping.standardField] = float(raw_value)
                except ValueError:
                    invalid_numeric_count += 1
                    cleaned_row[mapping.standardField] = None
                    cleaning_logs.append(
                        f"Invalid numeric value: {original_column}={raw_value} at row {row_index}"
                    )
            else:
                cleaned_row[mapping.standardField] = raw_value
        cleaned_rows.append(cleaned_row)

    return ExperimentAssetCleaningPreview(
        assetId=asset.asset_id,
        originalName=asset.original_name,
        rowCount=asset.row_count,
        cleanedRows=cleaned_rows,
        quality=ExperimentCleaningQualitySummary(
            missingValueCount=missing_value_count,
            invalidNumericCount=invalid_numeric_count,
            unmappedColumnCount=unmapped_column_count,
        ),
        cleaningLogs=cleaning_logs,
    )


def _is_numeric_standard_field(standard_field: str) -> bool:
    return standard_field != "timestamp"


def _detect_feature_anomalies(
    field: str,
    timestamped_values: list[tuple[str, str | float | None]],
) -> list[ExperimentFeatureAnomalyHint]:
    hints: list[ExperimentFeatureAnomalyHint] = []
    numeric_points = [
        (timestamp, float(value))
        for timestamp, value in timestamped_values
        if isinstance(value, (int, float))
    ]
    for index in range(1, len(numeric_points)):
        previous_timestamp, previous_value = numeric_points[index - 1]
        timestamp, value = numeric_points[index]
        absolute_change = abs(value - previous_value)
        relative_change = absolute_change / abs(previous_value) if previous_value else absolute_change
        if absolute_change >= 5 or relative_change >= 0.2:
            hints.append(
                ExperimentFeatureAnomalyHint(
                    field=field,
                    timestamp=timestamp,
                    value=value,
                    reason=(
                        f"Large step change from {previous_value} at {previous_timestamp} "
                        f"to {value}."
                    ),
                )
            )
    return hints


def _suggest_field_mapping(
    column_name: str,
    dictionary_entries: list[ExperimentFieldDictionaryEntry] | None = None,
) -> ExperimentFieldMappingSuggestion:
    normalized = _normalize_column_name(column_name)
    for entry in dictionary_entries or []:
        if _normalize_column_name(entry.sourceColumn) == normalized:
            return ExperimentFieldMappingSuggestion(
                originalColumn=column_name,
                standardField=entry.standardField,
                displayName=entry.displayName,
                unit=entry.unit,
                confidenceScore=0.9,
                requiresHumanConfirmation=False,
                reason="Matched enterprise field dictionary entry.",
            )

    known_fields: dict[str, tuple[str, str, str | None, str]] = {
        "time": ("timestamp", "时间戳", None, "常见时间列名称，可作为多设备时间轴对齐基准。"),
        "timestamp": ("timestamp", "时间戳", None, "标准时间戳字段，可作为多设备时间轴对齐基准。"),
        "datetime": ("timestamp", "时间戳", None, "常见日期时间列名称，可作为多设备时间轴对齐基准。"),
        "temp": ("temperature_c", "温度", "°C", "设备常用温度列缩写，默认按摄氏度处理。"),
        "temperature": ("temperature_c", "温度", "°C", "设备温度列，默认按摄氏度处理。"),
        "temperaturec": ("temperature_c", "温度", "°C", "字段包含温度与摄氏单位信息。"),
        "ph": ("ph", "pH", None, "酸碱度字段，常用于发酵和培养条件分析。"),
        "rpm": ("agitation_rpm", "搅拌转速", "rpm", "设备常用转速单位，可用于工艺条件建模。"),
        "agitationrpm": ("agitation_rpm", "搅拌转速", "rpm", "字段包含搅拌转速含义。"),
        "od600": ("optical_density_600", "OD600", None, "常见菌体生长光密度指标。"),
        "od": ("optical_density_600", "OD600", None, "疑似光密度指标，需要结合设备确认波长。"),
    }
    if normalized in known_fields:
        standard_field, display_name, unit, reason = known_fields[normalized]
        return ExperimentFieldMappingSuggestion(
            originalColumn=column_name,
            standardField=standard_field,
            displayName=display_name,
            unit=unit,
            confidenceScore=0.95 if normalized != "od" else 0.72,
            requiresHumanConfirmation=normalized == "od",
            reason=reason,
        )

    return ExperimentFieldMappingSuggestion(
        originalColumn=column_name,
        standardField=None,
        displayName="未识别字段",
        unit=None,
        confidenceScore=0.2,
        requiresHumanConfirmation=True,
        reason="当前规则库无法稳定识别该字段，需要人工确认或扩展企业字段词典。",
    )


def _normalize_column_name(column_name: str) -> str:
    return "".join(character.lower() for character in column_name if character.isalnum())


if "pytest" in sys.modules:
    experiment_file_service = ExperimentFileService()
else:
    _project_root = Path(__file__).resolve().parents[4]
    experiment_file_service = ExperimentFileService(SqliteExperimentDataStore(
        os.getenv(
            "EXPERIMENT_DATA_STORE_PATH",
            str(_project_root / "data" / "runtime" / "experiment-data.db"),
        ),
        os.getenv(
            "EXPERIMENT_OBJECT_STORE_PATH",
            str(_project_root / "data" / "runtime" / "experiment-objects"),
        ),
    ))
