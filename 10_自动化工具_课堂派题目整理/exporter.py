"""
导出模块
负责将题目导出为Excel和Word格式
"""
from typing import List
from pathlib import Path
from models import Question
import pandas as pd
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH


class Exporter:
    """
    导出器
    
    功能:
    1. 导出Excel文件
    2. 导出Word文件
    3. 去重处理
    """
    
    @staticmethod
    def export_to_excel(questions: List[Question], 
                       output_path: str,
                       include_raw_text: bool = True) -> bool:
        """
        导出为Excel文件
        
        参数:
            questions: 题目列表
            output_path: 输出路径
            include_raw_text: 是否包含原始文本
            
        返回:
            是否成功
        """
        try:
            if len(questions) == 0:
                return False
            
            # 去重
            questions = Exporter._deduplicate(questions)
            
            # 转换为DataFrame
            data = []
            for q in questions:
                row = {
                    "序号": q.question_index,
                    "题型": q.question_type,
                    "题干": q.question_text,
                    "选项": q.options,
                    "可见答案": q.visible_answer,
                    "解析": q.analysis,
                }
                if include_raw_text:
                    row["原始文本"] = q.raw_text
                data.append(row)
            
            df = pd.DataFrame(data)
            
            # 保存Excel
            output_file = Path(output_path)
            output_file.parent.mkdir(parents=True, exist_ok=True)
            
            with pd.ExcelWriter(output_file, engine='openpyxl') as writer:
                df.to_excel(writer, index=False, sheet_name='题目')
                
                # 获取workbook和worksheet
                workbook = writer.book
                worksheet = writer.sheets['题目']
                
                # 设置样式
                # 1. 第一行加粗
                for cell in worksheet[1]:
                    cell.font = Font(bold=True, size=11)
                    cell.fill = PatternFill(start_color="4472C4", 
                                           end_color="4472C4", 
                                           fill_type="solid")
                    cell.font = Font(bold=True, color="FFFFFF", size=11)
                
                # 2. 自动调整列宽
                for idx, col in enumerate(worksheet.columns, start=1):
                    max_length = 0
                    column_letter = get_column_letter(idx)
                    
                    for cell in col:
                        try:
                            if cell.value:
                                cell_length = len(str(cell.value))
                                if cell_length > max_length:
                                    max_length = cell_length
                        except:
                            pass
                    
                    # 设置列宽(最大50,最小10)
                    adjusted_width = min(max(max_length + 2, 10), 50)
                    worksheet.column_dimensions[column_letter].width = adjusted_width
                
                # 3. 文本自动换行
                for row in worksheet.iter_rows():
                    for cell in row:
                        cell.alignment = Alignment(wrap_text=True, vertical='top')
                
                # 4. 冻结首行
                worksheet.freeze_panes = 'A2'
            
            return True
            
        except Exception as e:
            print(f"导出Excel失败: {e}")
            return False
    
    @staticmethod
    def export_to_word(questions: List[Question],
                      output_path: str) -> bool:
        """
        导出为Word文件
        
        参数:
            questions: 题目列表
            output_path: 输出路径
            
        返回:
            是否成功
        """
        try:
            if len(questions) == 0:
                return False
            
            # 去重
            questions = Exporter._deduplicate(questions)
            
            # 创建Word文档
            doc = Document()
            
            # 标题
            title = doc.add_heading('课堂派题目整理', 0)
            title.alignment = WD_ALIGN_PARAGRAPH.CENTER
            
            # 添加每道题
            for q in questions:
                # 题目标题
                question_title = f"第 {q.question_index} 题"
                if q.question_type:
                    question_title += f"【{q.question_type}】"
                
                doc.add_heading(question_title, level=1)
                
                # 题干
                doc.add_paragraph()
                p = doc.add_paragraph()
                p.add_run("题干：").bold = True
                doc.add_paragraph(q.question_text)
                
                # 选项
                if q.options:
                    doc.add_paragraph()
                    p = doc.add_paragraph()
                    p.add_run("选项：").bold = True
                    for opt_line in q.options.split('\n'):
                        if opt_line.strip():
                            # 不使用样式，直接添加
                            doc.add_paragraph(f"  {opt_line.strip()}")
                
                # 答案
                doc.add_paragraph()
                p = doc.add_paragraph()
                p.add_run("答案：").bold = True
                if q.visible_answer:
                    p.add_run(q.visible_answer)
                else:
                    p.add_run("页面未显示")
                
                # 解析
                doc.add_paragraph()
                p = doc.add_paragraph()
                p.add_run("解析：").bold = True
                if q.analysis:
                    doc.add_paragraph(q.analysis)
                else:
                    p.add_run("页面未显示")
                
                # 分隔线
                doc.add_paragraph("─" * 50)
                doc.add_paragraph()
            
            # 保存文件
            output_file = Path(output_path)
            output_file.parent.mkdir(parents=True, exist_ok=True)
            doc.save(output_file)
            
            return True
            
        except Exception as e:
            import traceback
            error_detail = traceback.format_exc()
            print(f"导出Word失败: {e}")
            print(f"详细错误:\n{error_detail}")
            return False
    
    @staticmethod
    def _deduplicate(questions: List[Question]) -> List[Question]:
        """
        去重 - 根据题目序号和来源URL去重，避免误删
        
        参数:
            questions: 题目列表
            
        返回:
            去重后的题目列表
        """
        seen = set()
        deduplicated = []
        
        for q in questions:
            # 使用序号+来源URL作为唯一标识（同一页面的同一题号只保留一个）
            key = f"{q.source_url}_{q.question_index}"
            if key not in seen:
                seen.add(key)
                deduplicated.append(q)
        
        return deduplicated
