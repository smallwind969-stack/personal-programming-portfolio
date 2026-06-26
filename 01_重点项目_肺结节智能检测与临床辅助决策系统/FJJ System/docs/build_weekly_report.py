from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


OUT = Path(__file__).resolve().parent / "肺结节智能辅助诊断系统_本周工作汇报.docx"


def set_run_font(run, size=None, bold=False, color=None):
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size:
        run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_paragraph_font(paragraph, size=11, bold=False, color="000000"):
    for run in paragraph.runs:
        set_run_font(run, size=size, bold=bold, color=color)


def add_heading(doc, text, level=1):
    paragraph = doc.add_heading("", level=level)
    run = paragraph.add_run(text)
    if level == 1:
        set_run_font(run, size=16, bold=True, color="2E74B5")
        paragraph.paragraph_format.space_before = Pt(14)
        paragraph.paragraph_format.space_after = Pt(8)
    else:
        set_run_font(run, size=13, bold=True, color="1F4D78")
        paragraph.paragraph_format.space_before = Pt(10)
        paragraph.paragraph_format.space_after = Pt(6)
    return paragraph


def add_body(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.line_spacing = 1.15
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run(text)
    set_run_font(run, size=11)
    return paragraph


def add_bullet(doc, text):
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.line_spacing = 1.15
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(text)
    set_run_font(run, size=11)
    return paragraph


def shade_cell(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False, fill=None):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    set_run_font(run, size=10.5, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    if fill:
        shade_cell(cell, fill)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False
    for idx, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[idx], header, bold=True, fill="F2F4F7")
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            set_cell_text(cells[idx], value)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.49)
    section.footer_distance = Inches(0.49)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)

    for style_name in ["Heading 1", "Heading 2", "List Bullet"]:
        style = styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def build_report():
    doc = Document()
    configure_document(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("肺结节智能辅助诊断系统本周工作汇报")
    set_run_font(title_run, size=20, bold=True, color="0B2545")
    title.paragraph_format.space_after = Pt(4)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle.add_run("汇报主题：前端检测系统梳理、医疗 Agent M1 实现与 Qwen 8B 本地部署")
    set_run_font(subtitle_run, size=11, color="555555")
    subtitle.paragraph_format.space_after = Pt(14)

    add_heading(doc, "一、本周工作概述", 1)
    add_body(
        doc,
        "本周围绕肺结节自动检测与识别系统的工程化升级展开工作，重点从原有前端检测系统出发，完成了项目架构剖析、医疗辅助 Agent 的 M1 版本设计与实现、本地大模型 Qwen 8B 的部署联调，以及启动脚本和运行环境问题修复。当前系统已从“CT 输入后输出结节检测结果”推进到“基于模型输出生成初步随访建议”的基础流程。",
    )

    add_heading(doc, "二、项目架构梳理", 1)
    add_body(
        doc,
        "现有项目主体是一个基于 Flask、PyTorch 和前端页面的肺结节智能辅助诊断系统。系统接收 CT/DICOM 影像数据，调用深度学习模型完成结节检测、假阳性过滤和良恶性识别，并将结果返回到前端展示。",
    )
    add_table(
        doc,
        ["模块", "当前作用", "本周理解与结论"],
        [
            ["前端页面", "负责患者信息录入、影像上传、检测结果展示", "已重新整理页面结构，并新增 Agent 随访建议入口"],
            ["Flask 后端", "提供上传、检测、识别等接口", "新增 /agent/followup 接口，用于连接检测结果与大模型建议"],
            ["DiagnosisEngine", "串联 CT 读取、重采样、模型推理与结果整合", "当前可输出结节位置、概率、面积、良恶性概率等基础信息"],
            ["深度学习模型", "检测模型、假阳性过滤模型、良恶性分类模型", "目前未显式输出毛刺、分叶、胸膜牵拉等影像征象"],
            ["Agent 模块", "根据检测结果生成结构化随访建议", "M1 阶段已实现规则兜底 + 本地 Qwen 生成总结"],
        ],
    )

    add_heading(doc, "三、医疗辅助 Agent M1 实现", 1)
    add_body(
        doc,
        "M1 阶段的目标不是直接替代医生判断，而是先打通从模型检测结果到随访建议的基础闭环。系统会读取当前模型可得到的结节数量、最大结节面积、检测概率、良恶性概率等信息，再结合保守规则生成风险等级和建议，并调用本地 Qwen 8B 生成面向医生或患者的自然语言解释。",
    )
    add_bullet(doc, "新增 agent/llm_client.py：封装本地 Ollama API，默认调用 qwen3:8b。")
    add_bullet(doc, "新增 agent/followup_agent.py：实现风险分层、建议模板、缺失信息提醒和大模型总结。")
    add_bullet(doc, "新增 Flask 接口 /agent/followup：前端可直接请求随访建议。")
    add_bullet(doc, "前端新增“第三步：生成随访建议”区域，实现检测结果到 Agent 输出的页面闭环。")
    add_bullet(doc, "Agent 输出包含 case_summary、risk_level、followup_suggestion、recommendations、missing_information、llm_summary、safety_note 等字段。")

    add_heading(doc, "四、Qwen 8B 本地部署与联调", 1)
    add_body(
        doc,
        "考虑到医疗影像系统涉及隐私数据，本周采用本地部署方式接入大模型。结合本机 4060 显卡和 32GB 内存，选用 8B 规模模型作为 M1 阶段方案，兼顾可运行性、响应速度和中文医学表达能力。",
    )
    add_table(
        doc,
        ["事项", "完成情况"],
        [
            ["Ollama 安装", "已安装到项目目录下的 Ollama 文件夹，避免系统路径和权限问题"],
            ["模型拉取", "已成功拉取 qwen3:8b，模型大小约 5.2GB"],
            ["服务联调", "已验证 Ollama /api/generate 可返回中文结果"],
            ["Agent 联调", "已验证 /agent/followup 能调用本地 Qwen 并返回随访建议"],
            ["隐私设计", "当前调用链在本地运行，不依赖外部云端大模型接口"],
        ],
    )

    add_heading(doc, "五、运行环境与启动脚本修复", 1)
    add_body(
        doc,
        "在双击启动 bat 时发现报错，主要原因包括 Windows 批处理对中文绝对路径解析不稳定，以及当前 Python 环境缺少 Flask 相关依赖且 pip 损坏。为保证项目可以演示和复现，本周对启动方式做了工程化修复。",
    )
    add_bullet(doc, "将 bat 中的绝对路径改为相对路径，避免中文路径编码导致命令解析失败。")
    add_bullet(doc, "新增项目本地 vendor 目录，用于放置 Flask、flask-cors、itsdangerous、blinker 等依赖。")
    add_bullet(doc, "新增 install_missing_flask_deps.py，在系统 pip 不可用时仍能安装必要依赖。")
    add_bullet(doc, "启动脚本会自动启动 Ollama 服务，并运行 Flask 系统到 http://127.0.0.1:5001。")

    add_heading(doc, "六、当前能力边界", 1)
    add_body(
        doc,
        "当前 M1 版本已经能完成基础流程，但仍需要明确它的边界。现有模型输出主要是结节位置、检测概率、面积与良恶性概率，尚不能直接从代码输出毛刺、分叶、空泡征、胸膜牵拉等更细粒度影像征象。因此，Agent 目前不会伪造这些医学特征，而是把它们列入缺失信息，提示后续需要进一步扩展模型或引入人工标注/影像组学特征。",
    )

    add_heading(doc, "七、阶段成果", 1)
    add_bullet(doc, "完成了原肺结节检测识别系统的架构梳理，明确了前端、后端、模型推理和 Agent 模块的职责。")
    add_bullet(doc, "完成医疗辅助 Agent M1，初步实现从检测结果到随访建议的闭环。")
    add_bullet(doc, "完成 Qwen 8B 本地部署，验证本地大模型可用于生成中文医学辅助说明。")
    add_bullet(doc, "修复 bat 启动问题，提高项目演示稳定性。")
    add_bullet(doc, "明确当前系统能力边界，为后续知识库、影像征象提取和医学指南对齐打基础。")

    add_heading(doc, "八、下周计划", 1)
    add_table(
        doc,
        ["优先级", "计划内容", "预期产出"],
        [
            ["P0", "完善 Agent 输入结构，将检测结果、患者基础信息、历史随访信息统一成标准病例上下文", "更稳定的 Agent 入参格式"],
            ["P0", "建立肺结节随访知识库，优先整理 Fleischner 指南、国内肺结节诊疗共识等规则", "可检索、可引用的随访知识库"],
            ["P1", "增加 RAG 检索流程，让 Agent 的建议基于指南条文生成", "降低大模型幻觉，提高医学可解释性"],
            ["P1", "补充前端展示，包括风险等级、建议依据、缺失信息提醒", "更适合演示和答辩的交互页面"],
            ["P2", "探索毛刺、分叶等影像征象的提取方案", "为后续多模态医学 Agent 升级做准备"],
        ],
    )

    add_heading(doc, "九、向老师汇报时的重点表达", 1)
    add_body(
        doc,
        "本周工作的核心价值在于：项目不再只是一个单点的检测模型演示，而是开始向“医学影像检测 + 本地大模型 + 专业知识库 + 随访建议 Agent”的工程化系统演进。M1 版本已经打通基础链路，后续重点是提升医学依据、知识库质量和临床可解释性。",
    )

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build_report()
