from pathlib import Path
import re

from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle


ROOT = Path(__file__).resolve().parents[1]
SOURCES = [
    ROOT / "docs" / "final" / "肺结节检测Agent_大创结题答辩讲稿.md",
    ROOT / "docs" / "world_frontier" / "frontier_gap_and_sprint_protocol.md",
]


def font_path():
    candidates = [
        Path("C:/Windows/Fonts/msyh.ttc"),
        Path("C:/Windows/Fonts/simhei.ttf"),
        Path("C:/Windows/Fonts/simsun.ttc"),
    ]
    for path in candidates:
        if path.exists():
            return path
    return None


def clean_inline(text):
    text = re.sub(r"`([^`]+)`", r"\1", text)
    text = text.replace("|", "｜")
    return text.strip()


def parse_blocks(markdown):
    blocks = []
    lines = markdown.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line:
            i += 1
            continue
        if line.startswith("#"):
            level = len(line) - len(line.lstrip("#"))
            blocks.append(("heading", level, clean_inline(line[level:].strip())))
            i += 1
            continue
        if line.startswith("|") and i + 1 < len(lines) and lines[i + 1].startswith("|"):
            table = []
            while i < len(lines) and lines[i].startswith("|"):
                parts = [clean_inline(part) for part in lines[i].strip("|").split("|")]
                if not all(re.fullmatch(r"[-:\s]+", part) for part in parts):
                    table.append(parts)
                i += 1
            blocks.append(("table", table))
            continue
        if line.startswith("- "):
            items = []
            while i < len(lines) and lines[i].startswith("- "):
                items.append(clean_inline(lines[i][2:]))
                i += 1
            blocks.append(("bullets", items))
            continue
        para = [line]
        i += 1
        while i < len(lines) and lines[i].strip() and not lines[i].startswith(("#", "- ", "|")):
            para.append(lines[i].rstrip())
            i += 1
        blocks.append(("paragraph", clean_inline(" ".join(para))))
    return blocks


def set_run_font(run, size=11, bold=False):
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.bold = bold


def export_docx(blocks, out_path):
    doc = Document()
    styles = doc.styles
    styles["Normal"].font.name = "Microsoft YaHei"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    styles["Normal"].font.size = Pt(11)
    for block in blocks:
        if block[0] == "heading":
            _, level, text = block
            p = doc.add_heading(level=min(level, 3))
            run = p.add_run(text)
            set_run_font(run, size=18 if level == 1 else 14 if level == 2 else 12, bold=True)
        elif block[0] == "paragraph":
            p = doc.add_paragraph()
            run = p.add_run(block[1])
            set_run_font(run)
        elif block[0] == "bullets":
            for item in block[1]:
                p = doc.add_paragraph(style="List Bullet")
                run = p.add_run(item)
                set_run_font(run)
        elif block[0] == "table":
            rows = block[1]
            if not rows:
                continue
            table = doc.add_table(rows=len(rows), cols=max(len(r) for r in rows))
            table.style = "Table Grid"
            for r, row in enumerate(rows):
                for c, value in enumerate(row):
                    cell = table.cell(r, c)
                    run = cell.paragraphs[0].add_run(value)
                    set_run_font(run, size=9, bold=(r == 0))
    doc.save(out_path)


def export_pdf(blocks, out_path):
    font = font_path()
    font_name = "Helvetica"
    if font:
        pdfmetrics.registerFont(TTFont("CNFont", str(font)))
        font_name = "CNFont"
    styles = getSampleStyleSheet()
    body = ParagraphStyle("CNBody", parent=styles["BodyText"], fontName=font_name, fontSize=10.5, leading=16)
    h1 = ParagraphStyle("CNH1", parent=styles["Heading1"], fontName=font_name, fontSize=18, leading=24, spaceAfter=8)
    h2 = ParagraphStyle("CNH2", parent=styles["Heading2"], fontName=font_name, fontSize=14, leading=20, spaceAfter=6)
    h3 = ParagraphStyle("CNH3", parent=styles["Heading3"], fontName=font_name, fontSize=12, leading=18, spaceAfter=5)
    bullet = ParagraphStyle("CNBullet", parent=body, leftIndent=12, firstLineIndent=-8)
    story = []
    for block in blocks:
        if block[0] == "heading":
            _, level, text = block
            story.append(Paragraph(text, h1 if level == 1 else h2 if level == 2 else h3))
        elif block[0] == "paragraph":
            story.append(Paragraph(block[1], body))
            story.append(Spacer(1, 3))
        elif block[0] == "bullets":
            for item in block[1]:
                story.append(Paragraph(f"• {item}", bullet))
            story.append(Spacer(1, 4))
        elif block[0] == "table":
            rows = block[1]
            if not rows:
                continue
            data = [[Paragraph(cell, body) for cell in row] for row in rows]
            table = Table(data, repeatRows=1)
            table.setStyle(TableStyle([
                ("FONTNAME", (0, 0), (-1, -1), font_name),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EAF3F5")),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#B8C7D0")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 4),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
            ]))
            story.append(table)
            story.append(Spacer(1, 6))
    doc = SimpleDocTemplate(
        str(out_path),
        pagesize=A4,
        leftMargin=18 * mm,
        rightMargin=18 * mm,
        topMargin=18 * mm,
        bottomMargin=18 * mm,
    )
    doc.build(story)


def main():
    outputs = []
    for source in SOURCES:
        markdown = source.read_text(encoding="utf-8")
        blocks = parse_blocks(markdown)
        docx_path = source.with_suffix(".docx")
        pdf_path = source.with_suffix(".pdf")
        export_docx(blocks, docx_path)
        export_pdf(blocks, pdf_path)
        outputs.extend([docx_path, pdf_path])
    for path in outputs:
        print(path)


if __name__ == "__main__":
    main()
