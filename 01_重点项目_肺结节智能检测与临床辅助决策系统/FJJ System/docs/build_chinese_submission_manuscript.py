import os
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    Image as RLImage,
    ListFlowable,
    ListItem,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


BASE = Path(__file__).resolve().parent
OUT = BASE / "submission_cn"
FIG = OUT / "figures"
OUT.mkdir(exist_ok=True)
FIG.mkdir(exist_ok=True)

FONT_MS_YAHEI = r"C:\Windows\Fonts\msyh.ttc"
FONT_SIMSUN = r"C:\Windows\Fonts\simsun.ttc"
FONT_SIMHEI = r"C:\Windows\Fonts\simhei.ttf"


def font(size=28, bold=False):
    path = FONT_SIMHEI if bold else FONT_MS_YAHEI
    return ImageFont.truetype(path, size=size)


def draw_center(draw, box, text, fnt, fill=(24, 36, 48)):
    x1, y1, x2, y2 = box
    lines = text.split("\n")
    heights = []
    widths = []
    for line in lines:
        bbox = draw.textbbox((0, 0), line, font=fnt)
        widths.append(bbox[2] - bbox[0])
        heights.append(bbox[3] - bbox[1])
    total_h = sum(heights) + 8 * (len(lines) - 1)
    y = y1 + (y2 - y1 - total_h) / 2
    for line, w, h in zip(lines, widths, heights):
        draw.text((x1 + (x2 - x1 - w) / 2, y), line, font=fnt, fill=fill)
        y += h + 8


def arrow(draw, start, end, fill=(45, 95, 130), width=5):
    draw.line([start, end], fill=fill, width=width)
    ex, ey = end
    sx, sy = start
    if abs(ex - sx) >= abs(ey - sy):
        sign = 1 if ex > sx else -1
        pts = [(ex, ey), (ex - sign * 18, ey - 10), (ex - sign * 18, ey + 10)]
    else:
        sign = 1 if ey > sy else -1
        pts = [(ex, ey), (ex - 10, ey - sign * 18), (ex + 10, ey - sign * 18)]
    draw.polygon(pts, fill=fill)


def rounded_box(draw, box, text, fill, outline=(40, 72, 90), text_fill=(15, 23, 31)):
    draw.rounded_rectangle(box, radius=18, fill=fill, outline=outline, width=3)
    draw_center(draw, box, text, font(24, bold=True), text_fill)


def make_figures():
    # Figure 1: overall workflow
    img = Image.new("RGB", (1800, 980), "white")
    d = ImageDraw.Draw(img)
    d.text((60, 38), "图1  肺结节辅助诊断系统总体流程", font=font(38, True), fill=(12, 32, 45))
    boxes = [
        (80, 170, 360, 330, "DICOM/CT\n影像输入"),
        (450, 170, 730, 330, "1 mm重采样\nHU裁剪归一化"),
        (820, 170, 1100, 330, "候选结节\n高召回检测"),
        (1190, 170, 1470, 330, "假阳性精筛\n性质识别"),
        (1560, 170, 1760, 330, "医生\n复核"),
        (520, 560, 840, 740, "结构化病例上下文\n结节/风险/临床信息"),
        (980, 560, 1300, 740, "本地Agent\n风险分层与随访支持"),
        (1430, 560, 1710, 740, "报告草案\n缺失信息提醒"),
    ]
    palette = [(226, 244, 247), (232, 241, 255), (230, 249, 238), (255, 242, 225), (255, 235, 238)]
    for i, (x1, y1, x2, y2, t) in enumerate(boxes):
        rounded_box(d, (x1, y1, x2, y2), t, palette[i % len(palette)])
    for a, b in [((360, 250), (450, 250)), ((730, 250), (820, 250)), ((1100, 250), (1190, 250)), ((1470, 250), (1560, 250)), ((1660, 330), (680, 560)), ((840, 650), (980, 650)), ((1300, 650), (1430, 650))]:
        arrow(d, a, b)
    d.text((90, 860), "设计原则：模型只给出候选和风险提示；医生复核结果优先；Agent负责结构化总结、缺口提醒和安全边界表达。", font=font(28), fill=(52, 63, 72))
    path = FIG / "fig1_workflow.png"
    img.save(path, quality=95)

    # Figure 2: cascaded model
    img = Image.new("RGB", (1800, 980), "white")
    d = ImageDraw.Draw(img)
    d.text((60, 38), "图2  三级级联深度学习模型框架", font=font(38, True), fill=(12, 32, 45))
    boxes = [
        (90, 210, 430, 390, "AttentionVNet\n三维候选生成\n概率图/连通域"),
        (570, 210, 910, 390, "3D ResNet18-CBAM\n假阳性精筛\n候选保留/剔除"),
        (1050, 210, 1390, 390, "Dual-Branch ResNet\n局部+上下文融合\n良恶性倾向"),
        (1530, 210, 1740, 390, "输出\n位置/概率/风险"),
    ]
    for i, (x1, y1, x2, y2, t) in enumerate(boxes):
        rounded_box(d, (x1, y1, x2, y2), t, [(229, 245, 250), (240, 244, 255), (235, 249, 235), (255, 244, 224)][i])
    for a, b in [((430, 300), (570, 300)), ((910, 300), (1050, 300)), ((1390, 300), (1530, 300))]:
        arrow(d, a, b)
    rounded_box(d, (220, 600, 680, 760), "优化目标一：\n宽松检测，降低漏检风险", (246, 251, 255))
    rounded_box(d, (720, 600, 1180, 760), "优化目标二：\n精筛误报，降低复核负担", (246, 255, 248))
    rounded_box(d, (1220, 600, 1680, 760), "优化目标三：\n结合上下文，提升风险判断稳定性", (255, 249, 239))
    path = FIG / "fig2_cascade.png"
    img.save(path, quality=95)

    # Figure 3: Agent safety
    img = Image.new("RGB", (1800, 980), "white")
    d = ImageDraw.Draw(img)
    d.text((60, 38), "图3  本地医疗Agent的安全约束与人机闭环", font=font(38, True), fill=(12, 32, 45))
    rounded_box(d, (90, 180, 430, 330), "模型输出\n结节位置/置信度\n良恶性概率", (232, 244, 255))
    rounded_box(d, (90, 440, 430, 590), "医生复核\n确认/排除/新增\n补充征象", (230, 249, 238))
    rounded_box(d, (550, 300, 950, 470), "病例上下文合并\n医生结果优先\n缺失信息显式保留", (255, 248, 230))
    rounded_box(d, (1080, 180, 1450, 330), "规则层\n风险分层\n随访建议模板", (246, 239, 255))
    rounded_box(d, (1080, 440, 1450, 590), "本地Qwen/Ollama\n语言组织\n不改写事实", (245, 245, 245))
    rounded_box(d, (1550, 300, 1750, 470), "输出\n报告草案\n安全声明", (255, 235, 238))
    for a, b in [((430, 255), (550, 350)), ((430, 515), (550, 420)), ((950, 350), (1080, 255)), ((950, 420), (1080, 515)), ((1450, 255), (1550, 350)), ((1450, 515), (1550, 420))]:
        arrow(d, a, b)
    d.text((110, 800), "关键约束：不生成缺失的毛刺、分叶、胸膜牵拉、结节类型等影像征象；不使用“确诊”等替代医生判断的表述。", font=font(28), fill=(52, 63, 72))
    path = FIG / "fig3_agent_safety.png"
    img.save(path, quality=95)

    # Figure 4: roadmap
    img = Image.new("RGB", (1800, 980), "white")
    d = ImageDraw.Draw(img)
    d.text((60, 38), "图4  投稿前验证路线与证据闭环", font=font(38, True), fill=(12, 32, 45))
    steps = [
        ("锁定数据集", "病例数/结节数\n训练验证测试划分"),
        ("检测评价", "FROC曲线\n不同FP/scan敏感度"),
        ("分类评价", "AUC/F1/校准曲线\n混淆矩阵"),
        ("消融实验", "无CBAM/无双分支\n无Agent安全层"),
        ("专家评估", "指南一致性\n医生复核时间"),
    ]
    x = 90
    for title, body in steps:
        rounded_box(d, (x, 260, x + 270, 470), f"{title}\n{body}", (236, 248, 255))
        if title != steps[-1][0]:
            arrow(d, (x + 270, 365), (x + 350, 365))
        x += 350
    d.text((110, 680), "本文PDF已按方法学论文预稿组织；真正投稿前需补齐上述证据，否则只能作为系统实现论文或大创结题材料。", font=font(30, True), fill=(160, 58, 38))
    path = FIG / "fig4_validation_roadmap.png"
    img.save(path, quality=95)


TITLE = "基于三级级联深度学习与本地医疗Agent的肺结节CT辅助诊断和随访支持系统研究"
AUTHORS = "作者：待补充    单位：待补充"


SECTIONS = [
    ("摘要", [
        "目的：针对胸部CT肺结节筛查中候选病灶发现难、假阳性负担高、良恶性判断依赖多源信息以及随访建议难以结构化的问题，构建一种融合三级级联深度学习、医生复核和本地医疗Agent的肺结节辅助诊断系统。",
        "方法：系统以DICOM胸部CT序列为输入，首先通过SimpleITK完成序列读取、1 mm各向同性重采样、HU值裁剪和归一化处理；随后采用AttentionVNet生成三维候选结节概率图，利用3D ResNet18-CBAM对候选区域进行假阳性精筛，并通过Dual-Branch ResNet融合局部结节细节与周围组织上下文进行良恶性倾向识别。系统前端提供CT切片浏览、窗宽窗位调节、候选结节叠加显示、医生确认/排除/新增结节和临床信息补充等功能。进一步地，本文设计了本地医疗Agent，将模型输出、医生复核结果和临床上下文合并，生成病例摘要、风险分层、缺失信息提示和随访支持文本。",
        "结果：原型系统已实现从CT上传、候选检测、性质识别、医生复核到Agent报告草案生成的闭环流程。项目内部技术记录显示，系统在既有测试条件下达到检测召回率96.12%、检测精确率78.5%、良恶性分类准确率92.3%，单病例处理时间约25 s；上述指标仍需在锁定数据集、明确划分方案和统计置信区间后重新验证。Agent模块能够在不伪造缺失医学征象的前提下输出结构化风险解释和安全声明，并将既往CT对比、结节类型、最大径、毛刺、分叶、胸膜牵拉等缺失项显式列出。",
        "结论：该系统从单点模型推理扩展为“AI检测-医生复核-临床补充-Agent建议”的人机协同工作流，为肺结节筛查和随访管理提供了可审阅、可追溯、可本地部署的技术路径。未来需进一步完成FROC评价、消融实验、外部验证和医生读片实验，以支撑临床推广和高水平期刊投稿。",
        "关键词：肺结节；胸部CT；深度学习；计算机辅助诊断；人机协同；本地大模型；随访管理",
    ]),
    ("1 引言", [
        "肺癌是严重威胁人群健康的恶性肿瘤之一，早筛早诊对改善患者预后具有重要意义。低剂量胸部CT已成为肺癌筛查和肺结节评估的重要影像手段，但一次CT检查通常包含数百张切片，微小结节、磨玻璃结节、血管旁结节和胸膜旁结节容易受到噪声、解剖结构和阅片疲劳影响而被漏检。中文综述指出，肺结节形态小、边界模糊、密度分布不均，同时放射科医生面临阅片量增加和优质医疗资源分布不均等问题，CAD系统因此具有现实意义[1]。",
        "肺结节管理并不止于检测。临床医生还需结合结节大小、密度、表面征象、增长情况、患者危险因素和既往影像进行综合判断。国内肺结节数据标注和质量控制专家共识对实性结节、亚实性结节、纯磨玻璃结节、部分实性结节等概念进行了规范[2]；2024年中国专家共识也强调，亚实性结节随访应关注结节大小、实性成分变化以及AI和人机MDT评估[3]。这些临床要求提示，肺结节AI系统若仅输出一个检测框或分类分数，难以真正融入诊疗流程。",
        "近年来，深度学习推动了肺结节检测、分割和分类任务的发展。LIDC-IDRI和LUNA16等公开数据集促进了算法验证，U-Net、3D CNN、ResNet和注意力机制等方法被广泛应用于医学影像任务[4-7]。然而，现有研究仍存在两个突出问题：其一，高召回检测与低假阳性之间存在矛盾，单一模型难以同时兼顾；其二，许多系统停留在模型推理层面，缺少医生复核、结构化病例上下文和随访建议生成的闭环设计。",
        "本文提出一种基于三级级联深度学习与本地医疗Agent的肺结节CT辅助诊断系统。系统采用“先宽松检测保召回，再精筛降误报，最后融合上下文判风险”的级联策略，并在上层引入医生复核和本地Agent。Agent并不替代医生诊断，而是基于结构化输入生成病例摘要、风险层级、缺失信息提醒和安全边界说明。本文的核心贡献包括：（1）构建三级级联肺结节检测与识别流程；（2）实现支持医生复核的人机协同Web工作站；（3）设计本地医疗Agent，形成从模型结果到随访支持的闭环；（4）明确系统能力边界和投稿前验证路线。",
    ]),
    ("2 相关研究", [
        "肺结节检测领域早期主要依赖阈值分割、形态学处理和手工特征分类。随着深度学习发展，卷积神经网络能够从CT影像中自动学习纹理、边界和三维空间特征，逐渐成为主流技术路线。中文研究综述认为，基于深度学习的CT肺结节检测算法显著改善了检测精度和效率，但仍需解决低剂量噪声、复杂解剖结构、微小结节和应用落地之间的脱节问题[8]。",
        "国际上，LIDC-IDRI提供了1018例胸部CT及多名胸部放射科医师标注，是肺结节CAD研究的重要数据基础[4]。LUNA16基于LIDC-IDRI建立了自动肺结节检测客观评估框架，并推动了FROC等检测指标的标准化使用[5]。这些公开资源为模型训练和性能比较提供了基础，但临床系统还需要解决医生审阅、结果追溯和随访管理等问题。",
        "在模型结构方面，U-Net及其三维扩展适合医学影像中兼顾上下文和定位的分割或候选生成任务[6]。3D深度监督网络等体数据方法进一步表明，三维卷积能够直接学习体素级空间关系[7]。对于肺结节假阳性筛除和良恶性分类，残差网络、注意力机制和多尺度上下文融合已成为常见思路。本文采用AttentionVNet、3D ResNet18-CBAM和Dual-Branch ResNet的级联组合，正是为了分别优化候选生成、误报控制和风险识别。",
        "在临床工作流方面，肺结节随访研究提示，结节直径、体积、平均CT值、表面征象和恶性概率等因素与结节增长相关[9]。最新精准管理共识也指出，AI已被用于高危人群筛选、风险预测、影像阅片和随访策略制定，但不同人群交叉验证不足、假阳性标记过多等问题仍然存在[10]。因此，本文将Agent设计为医生复核后的辅助总结模块，而非独立诊断模块。",
    ]),
    ("3 材料与方法", [
        "3.1 系统输入与预处理。系统支持DICOM序列和MetaImage格式输入。后端使用SimpleITK读取CT序列，并将原始体数据重采样至1 mm各向同性空间，以减少不同扫描层厚和像素间距带来的尺度差异。随后将HU值裁剪到-1200至600 HU的肺窗范围，并归一化到0至1。为了支持前端快速浏览，系统同步生成PNG切片，但模型推理仍基于三维体数据。",
        "3.2 候选结节生成。第一阶段采用AttentionVNet进行三维候选结节检测。网络包含编码-解码结构、残差卷积块、SE通道重标定和注意力门控跳跃连接。推理阶段以64×64×64体素块进行滑窗扫描，步长为32体素，并对重叠区域概率进行平均融合。候选区域通过概率阈值和连通域分析得到。该阶段阈值设置相对宽松，以降低漏检风险。",
        "3.3 假阳性精筛。第二阶段使用3D ResNet18-CBAM对候选区域进行真假结节判别。系统以候选中心为基准提取32×32×32三维块，送入残差网络。CBAM通过通道注意力和空间注意力增强与结节相关的特征响应。该阶段主要用于剔除血管截面、支气管壁、炎症影和胸膜邻近结构等常见误报。",
        "3.4 良恶性倾向识别。第三阶段采用Dual-Branch ResNet融合局部与全局上下文信息。局部分支输入32体素尺度的结节中心块，全局分支输入64体素尺度的周围组织区域并重采样到相同尺寸。两个分支提取的512维特征拼接后进入全连接层，输出良性或恶性倾向概率。该设计对应临床判断中既看结节本体，也看邻近组织关系的思路。",
        "3.5 医生复核工作站。前端采用HTML、CSS、JavaScript和Canvas实现医学工作站风格界面，后端采用Flask提供REST接口。医生可浏览CT切片、调节窗宽窗位、查看模型候选结节、确认或排除候选、手动画框新增漏检结节，并补充年龄、性别、吸烟史、既往CT对比、既往肿瘤史和症状等信息。医生复核结果被保存为结构化病例上下文。",
        "3.6 本地医疗Agent。Agent接收模型候选、医生复核结果和临床上下文，输出病例摘要、结节统计、高中低风险分层、结节级建议、缺失信息、汇报亮点和安全声明。规则层优先采用医生确认和医生新增结节；若缺少最大径、结节类型、毛刺、分叶、胸膜牵拉、既往CT对比等信息，则明确列入缺失项。可选语言模型通过Ollama调用本地Qwen 8B，仅用于语言组织，不改变结构化事实和风险规则。",
    ]),
    ("4 系统实现与图示", [
        "系统总体流程如图1所示。输入层完成DICOM/CT数据接收和预处理；模型层完成候选检测、假阳性精筛和性质识别；交互层支持医生复核和临床信息补充；Agent层输出结构化随访支持和报告草案。",
        "[[FIG:fig1_workflow.png]]",
        "三级级联模型结构如图2所示。AttentionVNet负责高召回候选生成，3D ResNet18-CBAM负责假阳性控制，Dual-Branch ResNet负责局部与上下文融合的良恶性倾向识别。",
        "[[FIG:fig2_cascade.png]]",
        "Agent安全闭环如图3所示。系统将模型输出与医生复核结果合并为病例上下文，风险规则和本地语言模型共同生成说明文本。该模块遵循医生复核优先、缺失信息显式保留、禁止替代诊断三项原则。",
        "[[FIG:fig3_agent_safety.png]]",
    ]),
    ("5 实验设计与阶段性结果", [
        "5.1 评价指标。肺结节检测任务建议采用FROC曲线、不同假阳性率下的敏感度、召回率、精确率和每例假阳性数等指标；假阳性精筛和良恶性分类任务建议采用准确率、敏感度、特异度、F1值、AUC和校准曲线等指标；系统层面建议记录单病例处理时间、显存占用、医生复核时间变化、结构化字段完整率和随访建议一致率。",
        "5.2 阶段性内部结果。项目技术文档记录的阶段性指标包括：检测召回率96.12%、检测精确率78.5%、良恶性分类准确率92.3%、单病例处理时间约25 s。这些结果说明原型系统具备继续研究和投稿打磨的基础。但由于当前尚未形成锁定测试集、病例数、结节数、置信区间和公开基线对照，本文将其定义为内部技术验证结果，而非最终临床验证结论。",
        "5.3 Agent模块功能验证。升级后的Agent能够基于病例上下文返回统计量、风险分层、缺失信息和安全声明。例如，当医生未补充毛刺、分叶、胸膜牵拉或结节类型时，系统不会自动生成这些征象，而是提示“仍需补充”。这一行为降低了大模型生成医学事实错误的风险。",
        "5.4 投稿前验证路线。为了达到可投稿标准，系统需要补齐锁定数据集、FROC、多基线比较、消融实验、外部验证和专家一致性评价。投稿前证据闭环如图4所示。",
        "[[FIG:fig4_validation_roadmap.png]]",
    ]),
    ("6 讨论", [
        "本文系统的主要价值在于将肺结节AI从单点推理扩展为可复核的流程工具。单纯的结节检测模型只能回答“哪里可能有结节”，而临床实际还需要回答“哪些结果可信”“哪些需要医生复核”“缺少哪些关键信息”“如何形成可追溯的随访建议”。本文通过三级级联模型、人机复核界面和本地Agent，将这些环节串联起来。",
        "三级级联结构有助于缓解高召回和低误报之间的矛盾。候选生成阶段保留更多可疑区域，避免因阈值过严造成漏检；假阳性精筛阶段减少血管、炎症和胸膜结构造成的误报；双分支分类阶段进一步整合结节局部形态和周围上下文。与单阶段模型相比，这种拆分方式更便于针对不同错误类型进行优化，也更符合工程调试和临床解释需要。",
        "Agent模块的设计重点不是追求生成文本的华丽程度，而是追求医学安全和事实可追溯。中文专家共识提示，肺结节管理正从单纯大小判断转向结合影像特征、生长动力学和临床因素的综合评估[10]。因此，Agent必须知道自己不知道什么。本文系统将缺失信息显式列出，并将医生复核置于模型结果之上，有助于减少自动化系统在医疗场景中的过度自信。",
        "本研究仍有明显局限。第一，当前指标来自项目内部技术记录，尚未完成严格外部验证。第二，系统尚不能自动输出毛刺、分叶、空泡征、胸膜牵拉、钙化等细粒度征象。第三，Agent随访建议尚未与专家标注或指南条款进行量化一致性比较。第四，系统是否能真正缩短医生阅片时间、减少漏检或改善随访一致性，还需要前瞻性或至少回顾性读片实验验证。第五，若使用真实临床数据，必须补充伦理审批、脱敏流程和数据可用性声明。",
    ]),
    ("7 结论", [
        "本文构建了一种基于三级级联深度学习与本地医疗Agent的肺结节CT辅助诊断和随访支持系统。系统实现了从DICOM输入、三维候选检测、假阳性精筛、良恶性倾向识别、医生复核到Agent结构化输出的闭环流程。阶段性结果表明，该系统具备进一步开展定量实验和投稿打磨的基础。未来工作应围绕锁定数据集评价、FROC曲线、消融实验、外部验证、专家一致性评估和医生工作负荷实验展开，使系统从大创原型进一步提升为可投稿、可复现、可审查的医学影像AI研究。",
    ]),
    ("伦理声明", [
        "本文当前版本为系统方法学与原型实现研究稿。若后续纳入真实临床CT数据进行训练、验证或测试，应获得所在机构伦理委员会批准或豁免，并完成数据脱敏和使用授权。系统输出仅作为辅助信息，不替代影像科、呼吸科或胸外科医生诊断和治疗决策。",
    ]),
    ("数据可用性声明", [
        "本文原型参考公开肺结节数据集和本地样例DICOM完成系统开发。正式投稿时应明确训练、验证和测试数据来源、病例数、结节数、划分方式及访问权限。若使用LIDC-IDRI或LUNA16，应遵守其数据使用协议；若使用院内数据，应提供受控访问或不可公开原因说明。",
    ]),
    ("作者贡献", [
        "概念设计：待补充；方法设计：待补充；软件实现：待补充；实验验证：待补充；论文撰写：待补充；指导与监督：待补充。正式投稿前应按CRediT贡献分类由全体作者确认。",
    ]),
    ("利益冲突声明", [
        "作者声明无已知利益冲突。正式投稿前需由全体作者确认。",
    ]),
    ("人工智能工具使用声明", [
        "本文在论文结构规划、语言组织、图示草案和文档排版过程中使用了AI工具辅助。所有科学结论、实验数据、参考文献、医学表述和最终稿件责任均应由作者和指导教师审核确认。AI工具未替代作者对数据真实性、伦理合规性和临床有效性的责任。",
    ]),
    ("参考文献", [
        "[1] 深度学习方法在肺结节诊断中的应用[J]. 激光与光电子学进展, 2021, 58(16):1600005.",
        "[2] 中华放射学杂志专家组. 胸部CT肺结节数据标注与质量控制专家共识(2018)[J]. 中华放射学杂志, 2018.",
        "[3] 肺结节诊治中国专家共识(2024年版). 呼吸界, 2024.",
        "[4] Armato S G, McLennan G, Bidaut L, et al. The Lung Image Database Consortium (LIDC) and Image Database Resource Initiative (IDRI): a completed reference database of lung nodules on CT scans[J]. Medical Physics, 2011, 38(2):915-931.",
        "[5] Setio A A A, Traverso A, de Bel T, et al. Validation, comparison, and combination of algorithms for automatic detection of pulmonary nodules in computed tomography images: the LUNA16 challenge[J]. Medical Image Analysis, 2017, 42:1-13.",
        "[6] Ronneberger O, Fischer P, Brox T. U-Net: Convolutional Networks for Biomedical Image Segmentation[C]//MICCAI 2015. Cham: Springer, 2015:234-241.",
        "[7] Dou Q, Yu L, Chen H, et al. 3D deeply supervised network for automated segmentation of volumetric medical images[J]. Medical Image Analysis, 2017, 41:40-54.",
        "[8] 深度学习驱动的CT影像肺结节检测：挑战与进展[J]. 中国医疗器械杂志, 2023, 47(2):163-172.",
        "[9] 基于人工智能随访预测肺结节增长的影响因素研究[J]. 中国全科医学, 2022.",
        "[10] 肺结节精准管理专家共识(2026年版). 肿瘤界, 2026.",
        "[11] MacMahon H, Naidich D P, Goo J M, et al. Guidelines for Management of Incidental Pulmonary Nodules Detected on CT Images: From the Fleischner Society 2017[J]. Radiology, 2017, 284(1):228-243.",
        "[12] Lakhani P, Sundaram B, Recht M P, et al. Integrating and adopting AI in the radiology workflow: a primer for standards and IHE profiles[J]. Radiology, 2024, 310(3):e232653.",
    ]),
]


def build_markdown():
    parts = [f"# {TITLE}\n", AUTHORS, ""]
    for heading, paragraphs in SECTIONS:
        parts.append(f"## {heading}")
        for p in paragraphs:
            parts.append(p)
        parts.append("")
    md = "\n\n".join(parts)
    path = OUT / "肺结节辅助诊断系统中文投稿预稿.md"
    path.write_text(md, encoding="utf-8")
    return path


def register_pdf_fonts():
    pdfmetrics.registerFont(TTFont("MSYH", FONT_MS_YAHEI))
    pdfmetrics.registerFont(TTFont("SimHei", FONT_SIMHEI))


def pdf_header_footer(canvas, doc):
    canvas.saveState()
    canvas.setFont("MSYH", 8)
    canvas.setFillColor(colors.HexColor("#66717A"))
    canvas.drawString(2 * cm, 1.1 * cm, "中文投稿预稿 | 肺结节CT辅助诊断与本地医疗Agent")
    canvas.drawRightString(A4[0] - 2 * cm, 1.1 * cm, f"第 {doc.page} 页")
    canvas.restoreState()


def figure_flowable(filename, caption):
    path = str(FIG / filename)
    return [
        RLImage(path, width=16.2 * cm, height=8.8 * cm),
        Paragraph(caption, STYLES["caption"]),
        Spacer(1, 0.25 * cm),
    ]


def build_pdf():
    register_pdf_fonts()
    pdf_path = OUT / "肺结节辅助诊断系统中文投稿预稿.pdf"
    doc = SimpleDocTemplate(
        str(pdf_path),
        pagesize=A4,
        rightMargin=2.0 * cm,
        leftMargin=2.0 * cm,
        topMargin=2.0 * cm,
        bottomMargin=1.8 * cm,
    )

    story = []
    story.append(Paragraph(TITLE, STYLES["title"]))
    story.append(Spacer(1, 0.25 * cm))
    story.append(Paragraph(AUTHORS, STYLES["meta"]))
    story.append(Spacer(1, 0.45 * cm))

    for heading, paragraphs in SECTIONS:
        story.append(Paragraph(heading, STYLES["h1"]))
        for p in paragraphs:
            if p.startswith("[[FIG:"):
                fname = p.replace("[[FIG:", "").replace("]]", "")
                captions = {
                    "fig1_workflow.png": "图1 系统总体流程：从DICOM输入到医生复核和Agent报告草案生成。",
                    "fig2_cascade.png": "图2 三级级联模型：候选生成、假阳性精筛和良恶性倾向识别。",
                    "fig3_agent_safety.png": "图3 Agent安全闭环：医生复核优先、缺失信息显式保留、语言模型不改写事实。",
                    "fig4_validation_roadmap.png": "图4 投稿前验证路线：从数据锁定到专家一致性评价。",
                }
                story.extend(figure_flowable(fname, captions[fname]))
            elif heading == "参考文献":
                story.append(Paragraph(p, STYLES["ref"]))
            else:
                story.append(Paragraph(p, STYLES["body"]))
        story.append(Spacer(1, 0.18 * cm))

    doc.build(story, onFirstPage=pdf_header_footer, onLaterPages=pdf_header_footer)
    return pdf_path


def build_docx():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    styles = doc.styles
    styles["Normal"].font.name = "Microsoft YaHei"
    styles["Normal"].font.size = Pt(10.5)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(TITLE)
    run.bold = True
    run.font.name = "SimHei"
    run.font.size = Pt(18)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run(AUTHORS).font.size = Pt(10)

    for heading, paragraphs in SECTIONS:
        h = doc.add_heading(heading, level=1)
        for p in paragraphs:
            if p.startswith("[[FIG:"):
                fname = p.replace("[[FIG:", "").replace("]]", "")
                doc.add_picture(str(FIG / fname), width=Inches(6.5))
            else:
                para = doc.add_paragraph(p)
                para.paragraph_format.first_line_indent = Pt(21)
                para.paragraph_format.line_spacing = 1.25

    path = OUT / "肺结节辅助诊断系统中文投稿预稿.docx"
    doc.save(path)
    return path


def init_styles():
    base = getSampleStyleSheet()
    return {
        "title": ParagraphStyle(
            "title",
            parent=base["Title"],
            fontName="SimHei",
            fontSize=18,
            leading=25,
            alignment=TA_CENTER,
            spaceAfter=8,
            textColor=colors.HexColor("#12303F"),
        ),
        "meta": ParagraphStyle(
            "meta",
            parent=base["Normal"],
            fontName="MSYH",
            fontSize=9.5,
            leading=14,
            alignment=TA_CENTER,
            textColor=colors.HexColor("#5C6670"),
        ),
        "h1": ParagraphStyle(
            "h1",
            parent=base["Heading1"],
            fontName="SimHei",
            fontSize=13,
            leading=18,
            spaceBefore=10,
            spaceAfter=5,
            textColor=colors.HexColor("#0F4C5C"),
        ),
        "body": ParagraphStyle(
            "body",
            parent=base["BodyText"],
            fontName="MSYH",
            fontSize=10.2,
            leading=17,
            firstLineIndent=21,
            alignment=TA_JUSTIFY,
            spaceAfter=5,
        ),
        "caption": ParagraphStyle(
            "caption",
            parent=base["BodyText"],
            fontName="MSYH",
            fontSize=9,
            leading=13,
            alignment=TA_CENTER,
            textColor=colors.HexColor("#4D5963"),
            spaceAfter=4,
        ),
        "ref": ParagraphStyle(
            "ref",
            parent=base["BodyText"],
            fontName="MSYH",
            fontSize=9,
            leading=13,
            leftIndent=10,
            firstLineIndent=-10,
            spaceAfter=3,
        ),
    }


STYLES = None


def main():
    global STYLES
    make_figures()
    STYLES = init_styles()
    md = build_markdown()
    docx = build_docx()
    pdf = build_pdf()
    print(md)
    print(docx)
    print(pdf)


if __name__ == "__main__":
    main()
